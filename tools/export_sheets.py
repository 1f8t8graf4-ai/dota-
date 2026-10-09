# Выгрузка всех текстов приложения в Excel-листы для проверки нейросетями.
# Запуск из корня dota-:  python3 tools/export_sheets.py 12.8 [id сцены …]   → папка переводы-12.8/ (Excel по фильмам + всё сразу + текстом/). Промт для проверки — переводы-<версия>/00 ПРОМТ — как проверять.md (пишется руками, копировать из прошлой версии).
import json,sys,os,re
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,PatternFill
from openpyxl.utils import get_column_letter
V=sys.argv[1];OUT=f'переводы-{V}'
S=json.load(open('data/scenes.json'))['SCENES']
TOP=json.load(open('data/topics.json'))['TOPICS']
SW=json.load(open('data/scenewords.json'))['SCENEWORDS']
G=json.load(open('data/glossary.json'))
KIND={'film':'фильм','series':'сериал','clip':'клип','interview':'интервью'}
show=lambda s:s.get('show') or s['title']
FILMS=[];[FILMS.append(show(s)) for s in S if show(s) not in FILMS]
ep=lambda i:f'{i+1:02d}'
de_=lambda s:s.get('lang')=='de'
J=lambda L:'\n'.join(x for x in L if x)

def rows_scenes(L):
    R=[]
    for s in L:
        R.append([s['id'],show(s),s['sub'],s.get('ep',''),KIND.get(s['kind'],s['kind']),'немецкий' if de_(s) else 'английский',
                  s.get('lvl'),s.get('lvlWhy',''),s.get('use'),s.get('useWhy',''),
                  J(f"{ep(i)}. {p['t']}" for i,p in enumerate(s['parts'])),
                  J(f"{ep(i)}. {p['prize'][0]} {p['prize'][1]}" for i,p in enumerate(s['parts']) if p.get('prize'))])
    return R
H_SC=['ID сцены','Фильм / человек','Описание сцены','Серия / подпись','Тип','Язык оригинала','Сложность 1–5','Почему такая сложность','Польза 1–5','Чему учит','Эпизоды','Значки-призы эпизодов']

def rows_ph(L):
    R=[]
    for s in L:
        de=de_(s)
        for i,p in enumerate(s['parts']):
            for n,f in enumerate(p['ph']):
                R.append([f"{s['id']}|{i+1}|{n+1}",show(s),s['sub'],f"{ep(i)}. {p['t']}",'пассивная (только понять)' if f.get('passive') else 'для заданий',
                    f['de'] if de else f['en'], f['ru'], f['en'] if de else f.get('de',''),
                    f.get('use',''), (f['lx'][0] if f.get('lx') else ''), (f['lx'][1] if f.get('lx') else ''),
                    f.get('more',''),
                    J(' — '.join(x for x in e if x) for e in f.get('ex') or []),
                    J(' — '.join(x for x in e if x) for e in f.get('exDe') or []),
                    f.get('fact',''), ' / '.join(f.get('trap') or []), f.get('gap',''), ' / '.join(f.get('gx') or []),
                    ', '.join(TOP.get(f"{s['id']}|{f['en']}",[])),f"{f['a']:.1f}–{f['b']:.1f}"])
    return R
H_PH=['ID','Фильм / человек','Сцена','Эпизод','Тип','Оригинал (EN; у немецких сцен — DE)','Перевод RU','Второй язык (DE; у немецких сцен — EN)','Когда применяется','Пример из жизни','Перевод примера','Подробнее (объяснение)','Ещё примеры (EN — RU — DE)','Примеры для немецкого (DE — RU)','Интересный факт','Ложные переводы (вопрос «что значит»)','Слово для пропуска (всегда в английском варианте фразы)','Ложные слова для пропуска','Темы','Время в сцене, с']

def rows_kw(L):
    R=[]
    for s in L:
        for i,p in enumerate(s['parts']):
            for n,f in enumerate(p['ph']):
                for k in f.get('kw') or []:
                    k=list(k)+['']*(6-len(k))
                    R.append([f"{s['id']}|{i+1}|{n+1}",show(s),f"{ep(i)}. {p['t']}",f['de'] if de_(s) else f['en'],k[0],k[1],k[2],k[3],k[4],k[5]])
    return R
H_KW=['ID фразы','Фильм / человек','Эпизод','Фраза','Как во фразе','Словарная форма','Перевод','Пример','Перевод примера','Когда применяется / заметка']

def rows_fact(L):
    R=[]
    for s in L:
        for i,p in enumerate(s['parts']):
            if p.get('fact'):R.append([f"{s['id']}|{i+1}",show(s),s['sub'],f"{ep(i)}. {p['t']}",'об эпизоде (награда в конце)','',p['fact']])
            for n,f in enumerate(p['ph']):
                if f.get('fact'):R.append([f"{s['id']}|{i+1}|{n+1}",show(s),s['sub'],f"{ep(i)}. {p['t']}",'к фразе',f['de'] if de_(s) else f['en'],f['fact']])
    return R
H_F=['ID','Фильм / человек','Сцена','Эпизод','Чей факт','Фраза','Факт']

def rows_sub(L):
    R=[]
    for s in L:
        de=de_(s)
        for i,p in enumerate(s['parts']):
            for r in s['subs']:
                if p['a']-0.3<=r[0]<p['b']:
                    cl=lambda x:str(x or '').replace('\\n',' ').replace('\n',' ')
                    o=cl(r[4] if de else r[2]);t=cl(r[2] if de else (r[4] if len(r)>4 else ''))
                    R.append([f"{s['id']}|{r[0]:.2f}",show(s),s['sub'],f"{ep(i)}. {p['t']}",f"{r[0]:.1f}",o,cl(r[3]),t])
    return R
H_SUB=['ID строки','Фильм / человек','Сцена','Эпизод','Время, с','Оригинал','Перевод RU','Второй язык (DE; у немецких сцен — EN)']

def rows_sw(L):
    R=[]
    for s in L:
        W=SW.get(s['id']) or {}
        for i,p in enumerate(s['parts']):
            e=W.get(str(i)) or {}
            for k in e.get('key',[]):R.append([s['id'],show(s),f"{ep(i)}. {p['t']}",'ключевое слово эпизода',k[0],k[1],k[2] if len(k)>2 else ''])
            for w,t in (e.get('gloss') or {}).items():R.append([s['id'],show(s),f"{ep(i)}. {p['t']}",'подсказка к слову в субтитрах',w,t,''])
    return R
H_SW=['ID сцены','Фильм / человек','Эпизод','Что это','Слово','Перевод RU','DE']

def rows_gl():
    return [['английский',w,t] for w,t in sorted(G['GLOSS_EN'].items())]+[['немецкий',w,t] for w,t in sorted(G['GLOSS_DE'].items())]
H_GL=['Язык','Слово','Перевод RU (показывается, когда нажимаешь на слово)']

INSTR=[
 ['Что это',f'Все тексты приложения «Языки по кино» (версия {V}): Telegram-приложение, где русскоязычные учат английский (и немного немецкий) по сценам из фильмов, сериалов и интервью. Смотришь кусок сцены → на полезных фразах всплывает объяснение → потом проверка → интервальное повторение.'],
 ['Кто учится','Русскоязычные подростки и молодые люди, многие — геймеры (в приложении есть Dota и CS). Уровень — от новичка до среднего.'],
 ['Листы','«Сцены» — описание сцен. «Фразы» — главное: каждая фраза для обучения со всем, что к ней показывается. «Важные слова» — карточки слов внутри фраз. «Факты» — интересные факты (награда за эпизод и к фразам). «Субтитры» — все реплики сцен с переводом. «Слова эпизодов» и «Глоссарий» — переводы слов, которые всплывают, когда нажимаешь на слово.'],
 ['ID','Колонка ID — адрес строки. В ответе всегда указывай ID, иначе правку не найти.'],
 ['1. Перевод','Точный и живой русский, как сказал бы человек, а не калька. Смысл — как в сцене. Мат и грубость оставлены специально — не смягчать, не цензурить.'],
 ['2. Когда применяется','Должно быть СИТУАЦИЕЙ («Когда …»), а не грамматикой. Проверь, что носители реально так говорят в этой ситуации. Из этих текстов строится вопрос «когда так говорят?» — ситуации разных фраз не должны быть похожи до неразличимости.'],
 ['3. Примеры','Пример из жизни и «ещё примеры» — естественный современный английский/немецкий, не книжный; перевод верный.'],
 ['4. Факты','ТОЛЬКО проверяемые. Если факт сомнительный, неточный или выдуманный — отметь «убрать» или дай исправленный вариант с источником. Лучше без факта, чем с ложным.'],
 ['5. Ложные переводы','Для вопроса «что значит фраза?». Каждый должен быть НЕВЕРНЫМ, но правдоподобным. Если какой-то ложный вариант на самом деле тоже правильный перевод — это ошибка.'],
 ['6. Слово для пропуска','В фразе пропускают это слово, ученик выбирает из вариантов. «Ложные слова для пропуска» не должны подходить по смыслу.'],
 ['7. Важные слова','Словарная форма, перевод, пример и его перевод — верные; пример естественный.'],
 ['8. Немецкий','Колонки DE — грамотный естественный немецкий.'],
 ['9. Субтитры','Опечатки, неверный перевод реплики, обрезанные строки.'],
 ['10. Глоссарий','Перевод слова — самый частый смысл в разговорной речи.'],
 ['Не трогать','Мат. Стиль, если он и так нормальный. Тексты песен не переводить и не проверять.'],
 ['Формат ответа','Только проблемы, таблицей: ID | лист | колонка | было | стало | почему | уверенность (высокая / средняя). Если всё хорошо — так и напиши. В конце: сколько строк проверено на каждом листе.'],
]

def sheet(wb,title,head,rows,widths,fill='1F2937'):
    ws=wb.create_sheet(title)
    ws.append(head)
    for r in rows:ws.append(['' if x is None else x for x in r])
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=fill);c.alignment=Alignment(wrap_text=True,vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
    for i,w in enumerate(widths):ws.column_dimensions[get_column_letter(i+1)].width=w
    ws.freeze_panes='B2' if len(head)>3 else 'A2'
    if rows:ws.auto_filter.ref=ws.dimensions
    return len(rows)

def book(path,L,glossary):
    wb=Workbook();ws=wb.active;ws.title='Как проверять'
    ws.append(['Пункт','Что делать'])
    for r in INSTR:ws.append(r)
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='7C2D12')
    for row in ws.iter_rows(min_row=2):
        row[0].font=Font(bold=True)
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.column_dimensions['A'].width=22;ws.column_dimensions['B'].width=120
    n={}
    n['Сцены']=sheet(wb,'Сцены',H_SC,rows_scenes(L),[22,18,40,18,10,12,10,40,9,40,30,30])
    n['Фразы']=sheet(wb,'Фразы',H_PH,rows_ph(L),[24,16,28,22,14,42,42,36,48,40,40,50,50,40,48,40,14,22,18,12],'14532D')
    n['Важные слова']=sheet(wb,'Важные слова',H_KW,rows_kw(L),[24,16,22,40,16,18,28,40,40,48])
    n['Факты']=sheet(wb,'Факты',H_F,rows_fact(L),[24,16,28,22,16,36,90],'7C2D12')
    n['Субтитры']=sheet(wb,'Субтитры',H_SUB,rows_sub(L),[26,16,28,22,8,60,60,50])
    sw=rows_sw(L)
    if sw:n['Слова эпизодов']=sheet(wb,'Слова эпизодов',H_SW,sw,[22,16,22,24,18,36,18])
    if glossary:n['Глоссарий']=sheet(wb,'Глоссарий',H_GL,rows_gl(),[12,24,60])
    wb.save(path);return n

# markdown-версия (если нейронка не читает Excel)
def md(path,L,title):
    o=[f'# {title} — тексты на проверку (версия {V})','','Как проверять — в файле «00 ПРОМТ — как проверять.md». Адрес строки — ID в скобках.','']
    for s in L:
        de=de_(s)
        o+=[f"## {show(s)} — {s['sub']} ({s.get('ep','')})",f"Сложность {s.get('lvl')}: {s.get('lvlWhy','')}  ",f"Польза {s.get('use')}: {s.get('useWhy','')}",'']
        for i,p in enumerate(s['parts']):
            o+=[f"### Эпизод {ep(i)}. {p['t']}  (ID `{s['id']}|{i+1}`)",'']
            if p.get('fact'):o+=[f"**Факт об эпизоде:** {p['fact']}",'']
            if p.get('prize'):o+=[f"Значок: {p['prize'][0]} {p['prize'][1]}",'']
            for n,f in enumerate(p['ph']):
                o.append(f"**{f['de'] if de else f['en']}**  (ID `{s['id']}|{i+1}|{n+1}`)"+(' — _пассивная, только понять_' if f.get('passive') else ''))
                o.append(f"- Перевод: {f['ru']}")
                o.append(f"- {'EN' if de else 'DE'}: {f['en'] if de else f.get('de','')}")
                if f.get('use'):o.append(f"- Когда применяется: {f['use']}")
                if f.get('lx'):o.append(f"- Пример из жизни: {f['lx'][0]} — {f['lx'][1]}")
                if f.get('more'):o.append(f"- Подробнее: {f['more']}")
                for e in f.get('ex') or []:o.append(f"- Ещё пример: {' — '.join(x for x in e if x)}")
                for e in f.get('exDe') or []:o.append(f"- Пример (DE): {' — '.join(x for x in e if x)}")
                if f.get('fact'):o.append(f"- Факт: {f['fact']}")
                for k in f.get('kw') or []:
                    k=list(k)+['']*(6-len(k));o.append(f"- Слово: **{k[1] or k[0]}** — {k[2]} | пример: {k[3]} — {k[4]} | когда: {k[5]}")
                if f.get('trap'):o.append(f"- Ложные переводы: {' / '.join(f['trap'])}")
                if f.get('gap'):o.append(f"- Пропуск (в английском варианте): «{f['gap']}»"+(f"; ложные слова: {' / '.join(f['gx'])}" if f.get('gx') else ''))
                o.append('')
            rows=[r for r in s['subs'] if p['a']-0.3<=r[0]<p['b']]
            o+=[f"Субтитры эпизода ({len(rows)}):",'',f"| время | {'DE' if de else 'EN'} | RU | {'EN' if de else 'DE'} |",'|---|---|---|---|']
            cl=lambda x:str(x or '').replace('\\n',' ').replace('\n',' ').replace('|','\\|')
            for r in rows:o.append(f"| {r[0]:.1f} | {cl(r[4] if de else r[2])} | {cl(r[3])} | {cl(r[2] if de else (r[4] if len(r)>4 else ''))} |")
            o.append('')
    open(path,'w').write('\n'.join(o))

ONLY=sys.argv[2:]   # 12.9: только новые сцены — python3 tools/export_sheets.py 12.9 sopranos-ralph → переводы-12.9/ с одним файлом
if ONLY:
    L=[s for s in S if s['id'] in ONLY];os.makedirs(OUT,exist_ok=True);name='Новое '+V
    print(name,book(f'{OUT}/{name}.xlsx',L,False));md(f'{OUT}/{name}.md',L,'Новое в '+V);sys.exit()
os.makedirs(f'{OUT}/текстом',exist_ok=True)
for old in os.listdir(OUT):
    if old.endswith('.xlsx'):os.remove(f'{OUT}/{old}')
tot=book(f'{OUT}/Все сцены {V}.xlsx',S,True)
print('всё:',tot)
for k,h in enumerate(FILMS):
    L=[s for s in S if show(s)==h];name=f"{k+1:02d} {h}"
    n=book(f'{OUT}/{name}.xlsx',L,False);md(f'{OUT}/текстом/{name}.md',L,h)
    print(name,n)
